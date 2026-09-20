"""DOCX 텍스트 추출 (자소서 / 포트폴리오).

docs/layer-rules.md 1절 / task-09

포트폴리오는 표로 프로젝트를 정리하는 경우가 흔해 표 안의 글자도 함께 읽는다.
"""

import io
from zipfile import BadZipFile

import docx
from docx.opc.exceptions import PackageNotFoundError
from docx.table import Table

from app.integrations.extract.base import (
    EXTRACT_ERROR_CORRUPTED,
    EXTRACT_ERROR_EMPTY,
    ExtractionResult,
    failed,
    truncate,
)
from app.shared.enums import DocumentExtractStatus


def _table_lines(table: Table) -> list[str]:
    """표 한 개를 줄 목록으로 편다. 입력: Table. 출력: 빈 줄을 뺀 줄 목록."""
    lines: list[str] = []
    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells]
        joined = " | ".join(cell for cell in cells if cell)
        if joined:
            lines.append(joined)
    return lines


def extract_docx_text(data: bytes, *, max_chars: int = 0) -> ExtractionResult:
    """DOCX 바이트에서 텍스트를 뽑는다.

    입력: 파일 바이트, max_chars(0 이면 자르지 않는다).
    출력: ExtractionResult.

    - 열 수 없으면 corrupted_file
    - 문단과 표에 글자가 하나도 없으면 empty_document
    - 길이 상한으로 잘렸으면 partial
    """
    try:
        document = docx.Document(io.BytesIO(data))
    # BadZipFile 은 OSError 계열이 아니다 — docx 는 zip 이라 깨진 파일이 여기로 온다.
    except (PackageNotFoundError, BadZipFile, KeyError, ValueError, OSError):
        return failed(EXTRACT_ERROR_CORRUPTED)

    lines: list[str] = [
        paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()
    ]
    table_count = 0
    for table in document.tables:
        table_count += 1
        lines.extend(_table_lines(table))

    if not lines:
        return failed(EXTRACT_ERROR_EMPTY)

    text, is_truncated = truncate("\n".join(lines).strip(), max_chars)
    return ExtractionResult(
        text=text,
        status=DocumentExtractStatus.PARTIAL if is_truncated else DocumentExtractStatus.SUCCEEDED,
        is_truncated=is_truncated,
        details={"paragraphCount": len(document.paragraphs), "tableCount": table_count},
    )
