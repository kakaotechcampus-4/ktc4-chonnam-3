"""문서 텍스트 추출 — 형식 판정, PDF, DOCX, TXT/MD.

spec/backend/features/documents.md / task-09

파일 크기 상한(10MB)은 integrations 가 아니라 service·router 가 본다.
여기서는 형식과 추출 결과만 검증한다.
"""

import io

import docx
import pytest

from app.integrations.extract.base import (
    EXTRACT_ERROR_CORRUPTED,
    EXTRACT_ERROR_EMPTY,
    EXTRACT_ERROR_NO_TEXT_LAYER,
    DocumentFormat,
    detect_format,
    truncate,
)
from app.integrations.extract.docx import extract_docx_text
from app.integrations.extract.pdf import extract_pdf_text
from app.integrations.extract.plain import extract_plain_text
from app.shared.enums import DocumentExtractStatus


def build_pdf(text: str) -> bytes:
    """텍스트 레이어가 있는 최소 PDF 를 만든다. 입력: 본문. 출력: PDF 바이트."""
    content = f"BT /F1 12 Tf 20 100 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<</Type/Catalog/Pages 2 0 R>>",
        b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
        b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 300 200]/Contents 4 0 R"
        b"/Resources<</Font<</F1 5 0 R>>>>>>",
        b"<</Length " + str(len(content)).encode() + b">>stream\n" + content + b"\nendstream",
        b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for index, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{index} 0 obj".encode() + body + b"endobj\n"
    xref_offset = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode() + b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer<</Size {len(objects) + 1}/Root 1 0 R>>\nstartxref\n{xref_offset}\n%%EOF".encode()
    )
    return bytes(out)


def build_blank_pdf(pages: int = 1) -> bytes:
    """텍스트 레이어가 없는 PDF. 스캔 이미지 PDF 와 같은 상태다."""
    from pypdf import PdfWriter

    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def build_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        added = document.add_table(rows=len(table), cols=len(table[0]))
        for row_index, row in enumerate(table):
            for column_index, cell in enumerate(row):
                added.cell(row_index, column_index).text = cell
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ── 형식 판정 ──────────────────────────────────────────


@pytest.mark.parametrize(
    ("filename", "expected"),
    [
        ("resume.pdf", DocumentFormat.PDF),
        ("PORTFOLIO.PDF", DocumentFormat.PDF),
        ("resume.docx", DocumentFormat.DOCX),
        ("notes.txt", DocumentFormat.TXT),
        ("README.md", DocumentFormat.MD),
    ],
)
def test_supported_formats(filename: str, expected: DocumentFormat) -> None:
    assert detect_format(filename) is expected


@pytest.mark.parametrize("filename", ["resume.hwp", "shot.png", "deck.pptx", "deck.ppt", "noext"])
def test_unsupported_formats_return_none(filename: str) -> None:
    assert detect_format(filename) is None


def test_mime_is_used_when_extension_is_missing() -> None:
    assert detect_format("upload", "application/pdf") is DocumentFormat.PDF


def test_extension_wins_over_mime() -> None:
    """업로드 MIME 은 비어 있거나 octet-stream 인 경우가 흔해 확장자를 우선한다."""
    assert detect_format("resume.pdf", "application/octet-stream") is DocumentFormat.PDF


def test_truncate() -> None:
    assert truncate("abcdef", 3) == ("abc", True)
    assert truncate("abc", 10) == ("abc", False)
    assert truncate("abc", 0) == ("abc", False)


# ── PDF ────────────────────────────────────────────────


def test_pdf_success() -> None:
    result = extract_pdf_text(build_pdf("Hello DEVON github.com/alice/repo"))

    assert result.status is DocumentExtractStatus.SUCCEEDED
    assert "Hello DEVON" in result.text
    assert result.is_truncated is False
    assert result.details == {"pageCount": 1, "extractedPages": 1}


def test_pdf_partial_when_truncated() -> None:
    result = extract_pdf_text(build_pdf("Hello DEVON"), max_chars=5)

    assert result.status is DocumentExtractStatus.PARTIAL
    assert result.is_truncated is True
    assert result.text == "Hello"


def test_pdf_without_text_layer_fails() -> None:
    """스캔 이미지 PDF 는 추출 실패로 처리한다. OCR 은 범위 밖이다."""
    result = extract_pdf_text(build_blank_pdf(pages=2))

    assert result.status is DocumentExtractStatus.FAILED
    assert result.error_code == EXTRACT_ERROR_NO_TEXT_LAYER
    assert result.succeeded is False


@pytest.mark.parametrize("data", [b"", b"not a pdf at all"])
def test_pdf_corrupted_fails(data: bytes) -> None:
    result = extract_pdf_text(data)

    assert result.status is DocumentExtractStatus.FAILED
    assert result.error_code == EXTRACT_ERROR_CORRUPTED


# ── DOCX ───────────────────────────────────────────────


def test_docx_success_includes_table_text() -> None:
    data = build_docx(
        ["포트폴리오", "   "],
        table=[["프로젝트", "DEVON"], ["링크", "https://github.com/a/b"]],
    )

    result = extract_docx_text(data)

    assert result.status is DocumentExtractStatus.SUCCEEDED
    assert "포트폴리오" in result.text
    assert "DEVON" in result.text
    assert "https://github.com/a/b" in result.text


def test_docx_partial_when_truncated() -> None:
    result = extract_docx_text(build_docx(["가나다라마바사"]), max_chars=3)

    assert result.status is DocumentExtractStatus.PARTIAL
    assert result.is_truncated is True


def test_docx_empty_fails() -> None:
    result = extract_docx_text(build_docx([]))

    assert result.error_code == EXTRACT_ERROR_EMPTY


def test_docx_corrupted_fails() -> None:
    """docx 는 zip 이라 깨진 파일이 BadZipFile 로 온다."""
    result = extract_docx_text(b"not a docx")

    assert result.error_code == EXTRACT_ERROR_CORRUPTED


# ── TXT / MD ───────────────────────────────────────────


def test_plain_utf8() -> None:
    result = extract_plain_text("안녕 DEVON".encode())

    assert result.status is DocumentExtractStatus.SUCCEEDED
    assert result.text == "안녕 DEVON"


def test_plain_cp949() -> None:
    """국내 사용자가 올리는 txt 는 cp949 인 경우가 드물지 않다."""
    result = extract_plain_text("안녕 DEVON".encode("cp949"))

    assert result.status is DocumentExtractStatus.SUCCEEDED
    assert result.text == "안녕 DEVON"


def test_plain_strips_bom() -> None:
    result = extract_plain_text("안녕".encode("utf-8-sig"))

    assert "﻿" not in result.text


def test_plain_undecodable_is_partial_not_failed() -> None:
    """내용을 통째로 버리는 것보다 깨진 글자라도 남기는 쪽이 낫다."""
    result = extract_plain_text(b"\xff\xfe\x00\x01\x99")

    assert result.status is DocumentExtractStatus.PARTIAL
    assert result.succeeded is True


def test_plain_empty_fails() -> None:
    assert extract_plain_text(b"   \n  ").error_code == EXTRACT_ERROR_EMPTY
