"""포트폴리오 링크 대상과 길이 축약의 보존 경계."""

import importlib
import io
import zipfile

import docx
import pytest
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pypdf import PdfReader, PdfWriter
from pypdf.annotations import Link
from pypdf.generic import (
    ArrayObject,
    DictionaryObject,
    NameObject,
    NullObject,
    NumberObject,
    TextStringObject,
)

from app.integrations.extract.base import DocumentFormat
from app.integrations.extract.docx import extract_docx_text
from app.integrations.extract.pdf import extract_pdf_text
from app.shared.enums import DocumentExtractStatus
from tests.integrations.test_extract import build_docx, build_pdf


def linked_pdf() -> bytes:
    writer = PdfWriter()
    writer.append(PdfReader(io.BytesIO(build_pdf("Repository link"))))
    writer.add_annotation(
        page_number=0,
        annotation=Link(rect=(0, 0, 100, 100), url="https://github.com/alice/pdf.git/tree/main"),
    )
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def linked_docx() -> bytes:
    document = docx.Document()
    paragraph = document.add_paragraph("Portfolio ")
    relationship = paragraph.part.relate_to(
        "https://github.com/alice/docx/issues/2", RELATIONSHIP_TYPE.HYPERLINK, is_external=True
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "Repository link"
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.mark.parametrize(
    ("extractor", "data", "expected"),
    [
        (extract_pdf_text, linked_pdf, "https://github.com/alice/pdf.git/tree/main"),
        (extract_docx_text, linked_docx, "https://github.com/alice/docx/issues/2"),
    ],
)
def test_hyperlink_target_is_retained_when_display_text_is_not_a_url(extractor, data, expected):
    result = extractor(data())
    assert expected in getattr(result, "hyperlinks", ()), "Hidden hyperlink target was lost"


def test_malformed_pdf_annotations_preserve_readable_text_as_partial():
    writer = PdfWriter()
    writer.append(PdfReader(io.BytesIO(build_pdf("Readable page"))))
    writer.pages[0][NameObject("/Annots")] = NullObject()
    buffer = io.BytesIO()
    writer.write(buffer)

    result = extract_pdf_text(buffer.getvalue())

    assert result.text == "Readable page"
    assert result.status is DocumentExtractStatus.PARTIAL


@pytest.mark.parametrize("bad_action", [NumberObject(3), TextStringObject("broken")])
def test_malformed_pdf_action_preserves_readable_text_as_partial(bad_action):
    writer = PdfWriter()
    writer.append(PdfReader(io.BytesIO(build_pdf("Readable page"))))
    writer.pages[0][NameObject("/Annots")] = ArrayObject(
        [DictionaryObject({NameObject("/A"): bad_action})]
    )
    buffer = io.BytesIO()
    writer.write(buffer)

    result = extract_pdf_text(buffer.getvalue())

    assert result.text == "Readable page"
    assert result.status is DocumentExtractStatus.PARTIAL


def test_pdf_indirect_hyperlink_target_is_retained():
    writer = PdfWriter()
    writer.append(PdfReader(io.BytesIO(build_pdf("Repository link"))))
    target = writer._add_object(TextStringObject("https://github.com/alice/indirect-uri"))
    action = DictionaryObject({NameObject("/S"): NameObject("/URI"), NameObject("/URI"): target})
    writer.pages[0][NameObject("/Annots")] = ArrayObject(
        [DictionaryObject({NameObject("/A"): action})]
    )
    buffer = io.BytesIO()
    writer.write(buffer)

    result, urls = process(buffer.getvalue(), DocumentFormat.PDF)

    assert result.status is DocumentExtractStatus.SUCCEEDED
    assert urls == ["alice/indirect-uri"]


@pytest.mark.parametrize(
    "document_xml",
    [
        b"<wrong/>",
        b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>',
    ],
    ids=["wrong_root", "missing_body"],
)
def test_invalid_docx_structure_is_an_expected_extraction_failure(document_xml):
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(build_docx(["Original text"]))) as original:
        with zipfile.ZipFile(buffer, "w") as modified:
            for info in original.infolist():
                data = document_xml if info.filename == "word/document.xml" else original.read(info)
                modified.writestr(info, data)

    result = extract_docx_text(buffer.getvalue())

    assert result.status is DocumentExtractStatus.FAILED
    assert result.error_code == "corrupted_file"


def test_unused_docx_hyperlink_relationship_is_not_a_portfolio_mention():
    document = docx.Document()
    document.add_paragraph("No repository mentioned")
    document.part.relate_to(
        "https://github.com/alice/deleted-project", RELATIONSHIP_TYPE.HYPERLINK, is_external=True
    )
    buffer = io.BytesIO()
    document.save(buffer)

    result, urls = process(buffer.getvalue(), DocumentFormat.DOCX)

    assert result.status is DocumentExtractStatus.SUCCEEDED
    assert urls == []


def process(data: bytes, document_format=DocumentFormat.TXT, max_chars=120):
    assert importlib.util.find_spec("app.features.documents"), "Document processing is missing"
    module = importlib.import_module("app.features.documents.processing")
    return module.process_document(data, document_format, max_chars=max_chars)


def test_truncation_keeps_end_links_and_prioritizes_url_context_and_project_heading():
    source = (
        "Unrelated biography.\n" * 30
        + "# Project DEVON\nBuilt an interview service.\n"
        + "Unrelated employment.\n" * 30
        + "Repository implementation:\nhttps://github.com/alice/end.git/tree/main?x=1\n"
    )
    result, urls = process(source.encode(), max_chars=150)

    assert urls == ["alice/end"]
    assert len(result.text) <= 150
    assert "# Project DEVON" in result.text
    assert "Repository implementation:" in result.text
    assert "https://github.com/alice/end.git/tree/main?x=1" in result.text
    assert result.text.index("# Project DEVON") < result.text.index("Repository implementation:")
    assert result.status is DocumentExtractStatus.PARTIAL
    assert result.is_truncated is True
    assert process(source.encode(), max_chars=150) == (result, urls)


def test_nul_cleanup_is_partial_and_does_not_lose_normalized_urls():
    result, urls = process(b"Hello\x00 world github.com/alice/project\x00")
    assert result.text == "Hello world github.com/alice/project"
    assert result.status is DocumentExtractStatus.PARTIAL
    assert result.is_truncated is False
    assert urls == ["alice/project"]


def test_nul_only_document_is_failed_after_cleanup():
    result, urls = process(b"\x00\x00")
    assert result.status is DocumentExtractStatus.FAILED
    assert result.error_code == "empty_document"
    assert urls == []


@pytest.mark.parametrize("max_chars", [1, 10, 120])
def test_all_urls_survive_even_when_text_budget_cannot_fit_them(max_chars):
    source = "github.com/alice/first\n" + "Filler.\n" * 100 + "github.com/alice/last"
    result, urls = process(source.encode(), max_chars=max_chars)
    assert urls == ["alice/first", "alice/last"]
    assert len(result.text) <= max_chars


@pytest.mark.parametrize(
    ("document_format", "data", "expected"),
    [
        (DocumentFormat.PDF, linked_pdf, ["alice/pdf"]),
        (DocumentFormat.DOCX, linked_docx, ["alice/docx"]),
    ],
)
def test_processing_includes_hidden_hyperlinks(document_format, data, expected):
    result, urls = process(data(), document_format, max_chars=120)
    assert urls == expected
    assert result.status is DocumentExtractStatus.SUCCEEDED


def test_nonpositive_text_budget_is_rejected():
    with pytest.raises(ValueError):
        process(b"text", max_chars=0)
